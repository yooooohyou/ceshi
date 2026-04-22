"""
嵌入组件（Embed）路由

- POST /doc_editor/embeds                 创建组件并返回可插入 HTML 的标记
- GET  /doc_editor/embeds/{embed_id}      获取组件详情（payload）
- PUT  /doc_editor/embeds/{embed_id}      更新组件
- DELETE /doc_editor/embeds/{embed_id}    软删除组件
- GET  /doc_editor/embeds/{embed_id}/go   URL 跳转（302 → url）
- POST /doc_editor/embeds/scan            从 HTML 扫描所有 embed 标记
- GET  /doc_editor/embeds/by_record/{record_id}  列出某个文档下所有组件
- POST /doc_editor/embeds/xlsx2html_config  解析上传的xlsx，返回结构化配置信息
- GET  /doc_editor/embeds/test/html       【测试】预览表格前10行 + 全部数据链接
- GET  /doc_editor/embeds/test/html/full  【测试】全部表格数据
- POST /doc_editor/embeds/test/docx       【测试】DOCX 表格替换渲染
"""
import csv
import html as html_lib
import logging
import math
import os
from typing import Any, Dict, List, Optional

import tempfile

from fastapi import APIRouter, Body, Path, Query
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse

from app.db.database import (
    delete_embed_component,
    get_embed_component,
    get_db_connection,
    insert_embed_component,
    list_embed_components_by_record,
    update_embed_component,
)
from app.models.schemas import unified_response
from app.utils.embed_marker import (
    DISPLAY_BLOCK,
    DISPLAY_INLINE,
    TYPE_TABLE,
    EmbedSpec,
    build_embed_marker,
    parse_embed_markers_from_html,
    spec_from_db_row,
    spec_to_db_row,
)

# 导入即完成渲染器注册
import app.utils.embed_table_renderer  # noqa: F401

router = APIRouter()
logger = logging.getLogger(__name__)


# ─── 创建 ─────────────────────────────────────────────────────────────────────

@router.post("/embeds", summary="创建嵌入组件并返回可插入HTML的标记")
async def create_embed(
    data: Dict[str, Any] = Body(..., description="组件业务数据(JSON)"),
    embed_type: str = Body("custom", description="组件类型：text/table/image/chart/reference/custom"),
    title: str = Body("", description="标题（显示用）"),
    url: Optional[str] = Body(None, description="跳转URL；不传则自动生成 /doc_editor/embeds/{id}/go"),
    display: Optional[str] = Body(None, description="inline / block；不传按类型自动判定"),
    version: int = Body(1, description="payload结构版本"),
    record_id: Optional[int] = Body(None, description="关联文档记录ID"),
    node_id: Optional[int] = Body(None, description="关联节点ID"),
):
    try:
        spec, snippet = build_embed_marker(
            data=data,
            embed_type=embed_type,
            title=title,
            url=url,
            display=display,
            version=version,
            record_id=record_id,
            node_id=node_id,
            url_builder=lambda eid: f"/doc_editor/embeds/{eid}/go",
        )
    except ValueError as e:
        return unified_response(400, str(e))

    try:
        insert_embed_component(spec_to_db_row(spec))
    except Exception as e:
        logger.error(f"create_embed: 入库失败 err={e}")
        return unified_response(500, f"入库失败：{str(e)}")

    return unified_response(200, "创建成功", {
        "embed_id": spec.embed_id,
        "url":      spec.url,
        "display":  spec.display,
        "html":     snippet,
        "spec":     spec.to_dict(),
    })


# ─── 查询 ─────────────────────────────────────────────────────────────────────

@router.get("/embeds/{embed_id}", summary="查询嵌入组件详情")
async def get_embed(embed_id: str = Path(..., description="embed_id，如 EMB_xxxxxxxxxxxx")):
    row = get_embed_component(embed_id)
    if not row:
        return unified_response(404, f"未找到 embed_id={embed_id}")
    spec = spec_from_db_row(row)
    return unified_response(200, "查询成功", spec.to_dict())


# ─── 更新 ─────────────────────────────────────────────────────────────────────

@router.put("/embeds/{embed_id}", summary="更新嵌入组件")
async def update_embed(
    embed_id: str = Path(...),
    data: Optional[Dict[str, Any]] = Body(None, description="新的业务数据；为空时不更新payload"),
    title: Optional[str] = Body(None),
    url: Optional[str] = Body(None),
    display: Optional[str] = Body(None),
    version: Optional[int] = Body(None),
):
    row = get_embed_component(embed_id)
    if not row:
        return unified_response(404, f"未找到 embed_id={embed_id}")

    spec = spec_from_db_row(row)
    if data is not None:
        spec.payload = data
    if title is not None:
        spec.title = title
    if url is not None:
        spec.url = url
    if display in (DISPLAY_INLINE, DISPLAY_BLOCK):
        spec.display = display
    if version is not None:
        spec.version = int(version)

    ok = update_embed_component(embed_id, spec_to_db_row(spec))
    if not ok:
        return unified_response(500, "更新失败")
    return unified_response(200, "更新成功", spec.to_dict())


# ─── 删除 ─────────────────────────────────────────────────────────────────────

@router.delete("/embeds/{embed_id}", summary="删除嵌入组件（软删除）")
async def delete_embed(embed_id: str = Path(...)):
    ok = delete_embed_component(embed_id)
    if not ok:
        return unified_response(404, f"未找到或已删除：{embed_id}")
    return unified_response(200, "删除成功", {"embed_id": embed_id})


# ─── URL 跳转 ─────────────────────────────────────────────────────────────────

@router.get("/embeds/{embed_id}/go", summary="根据embed_id跳转到组件URL")
async def go_embed(embed_id: str = Path(...)):
    row = get_embed_component(embed_id)
    if not row:
        return unified_response(404, f"未找到 embed_id={embed_id}")

    target = row.get("url")

    # 修复：如果 target 为空，或者 target 指向了接口自身，则不进行跳转，直接返回数据
    if not target or target.endswith(f"/embeds/{embed_id}/go"):
        spec = spec_from_db_row(row)
        return unified_response(200, "该组件无跳转URL", spec.to_dict())

    return RedirectResponse(url=target, status_code=302)


# ─── 扫描 HTML ────────────────────────────────────────────────────────────────

@router.post("/embeds/scan", summary="从HTML文本中扫描所有嵌入组件标记")
async def scan_embeds(html_text: str = Body(..., embed=True, description="HTML文本")):
    markers = parse_embed_markers_from_html(html_text)
    return unified_response(200, "扫描成功", {
        "count":   len(markers),
        "markers": markers,
    })


# ─── 列表 ─────────────────────────────────────────────────────────────────────

@router.get("/embeds/by_record/{record_id}", summary="列出某文档下所有嵌入组件")
async def list_by_record(record_id: int = Path(..., gt=0)):
    rows = list_embed_components_by_record(record_id)
    return unified_response(200, "查询成功", {
        "record_id": record_id,
        "count":     len(rows),
        "items":     rows,
    })


# ─── xlsx2html_config ─────────────────────────────────────────────────────────

def _resolve_xlsx_path(file_path: str) -> Optional[str]:
    """将 split_uploads 返回的 file_path 解析为本地绝对路径"""
    from urllib.parse import urlparse
    from app.core.config import UPLOAD_DIR

    if os.path.isfile(file_path):
        return file_path

    parsed = urlparse(file_path)
    raw = parsed.path if parsed.scheme in ("http", "https") else file_path
    basename = os.path.basename(raw.rstrip("/").rstrip("\\"))
    if basename:
        candidate = os.path.join(UPLOAD_DIR, basename)
        if os.path.isfile(candidate):
            return candidate

    return None


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return ""
        if v == int(v):
            return str(int(v))
    return str(v)


def _read_xlsx_sheets(path: str, max_rows: int = 10) -> List[Dict]:
    import pandas as pd
    with pd.ExcelFile(path, engine="openpyxl") as xl:
        sheets = []
        for sheet_name in xl.sheet_names:
            df = pd.read_excel(xl, sheet_name=sheet_name, header=None, nrows=max_rows)
            df = df.where(df.notna(), other=None)
            rows = [
                {"index": i + 1, "data": [{"index": j + 1, "value": _cell_to_str(v)} for j, v in enumerate(row)]}
                for i, row in enumerate(df.values.tolist())
            ]
            sheets.append({"sheetName": str(sheet_name), "top10Rows": rows})
    return sheets


def _find_xlsx_by_filename(file_name: str) -> Optional[str]:
    """在 UPLOAD_DIR 中按文件名查找 xlsx，支持直接命中和 {sign}_{name} 格式"""
    from app.core.config import UPLOAD_DIR

    # 只取 basename，防止路径穿越（如 file_name="/etc/passwd" 或含 ../）
    safe_name = os.path.basename(file_name)
    if not safe_name:
        return None

    upload_abs = os.path.abspath(UPLOAD_DIR)

    direct = os.path.join(upload_abs, safe_name)
    if os.path.isfile(direct):
        return direct

    try:
        suffix = "_" + safe_name
        matched = [
            fn for fn in os.listdir(upload_abs)
            if fn == safe_name or fn.endswith(suffix)
        ]
        if matched:
            matched.sort(key=lambda fn: os.path.getmtime(os.path.join(upload_abs, fn)), reverse=True)
            return os.path.join(upload_abs, matched[0])
    except Exception:
        pass

    return None


@router.post("/embeds/xlsx2html_config", summary="解析上传的xlsx文件，返回结构化配置信息")
async def xlsx2html_config(
    file_path: str = Body(..., description="split_uploads 返回的 file_path 或 real_file_path"),
    file_name: Optional[str] = Body(None, description="原始文件名，不传则从路径中提取"),
):
    actual_path = _resolve_xlsx_path(file_path)
    if not actual_path:
        return unified_response(400, f"文件不存在或路径无效：{file_path}")

    if not actual_path.lower().endswith(".xlsx"):
        return unified_response(400, "仅支持 xlsx 格式文件")

    try:
        sheets = _read_xlsx_sheets(actual_path)
    except Exception as e:
        logger.error(f"xlsx2html_config: 读取失败 path={actual_path} err={e}")
        return unified_response(500, f"解析 xlsx 失败：{str(e)}")

    resolved_name = file_name or os.path.basename(actual_path)

    return unified_response(200, "解析成功", {
        "fileName": resolved_name,
        "sheets": sheets,
    })


# ─── xlsx2html ─────────────────────────────────────────────────────────────────

@router.post(
    "/embeds/xlsx2html",
    summary="根据选定的 Sheet 和表头行生成 HTML 预览（前10行 + 全部数据链接），完整数据存入 DB 供 merge_docx 写入",
    response_class=HTMLResponse,
)
async def xlsx2html(
    fileName: str = Body(..., description="xlsx2html_config 返回的 fileName"),
    sheets: List[Dict[str, Any]] = Body(..., description="选定的 Sheet 列表，每项含 sheetName 和 headerRow（{index, data}，index 为 xlsx2html_config 返回的行序号，data 为选用的列名列表，可调整顺序或减少列）"),
    record_id: Optional[int] = Body(None, description="关联文档记录ID，传入后 merge_docx 可将完整数据写入 DOCX"),
    preview_rows: int = Body(10, ge=1, le=100, description="预览行数，默认10"),
):
    actual_path = _find_xlsx_by_filename(fileName)
    if not actual_path:
        return HTMLResponse(
            content=f"<!DOCTYPE html><html><body><h3>未找到文件：{html_lib.escape(fileName)}</h3></body></html>",
            status_code=404,
        )

    # 用磁盘上的实际文件名（new_filename）查 DB，取回用户上传时的原始文件名作为标题展示
    from app.db.database import get_original_filename_by_new_filename
    _new_fname = os.path.basename(actual_path)
    display_name = get_original_filename_by_new_filename(_new_fname) or fileName

    import pandas as pd

    parts: List[str] = []
    try:
        xl = pd.ExcelFile(actual_path, engine="openpyxl")
    except Exception as e:
        logger.error(f"xlsx2html: 打开文件失败 path={actual_path} err={e}")
        return HTMLResponse(
            content=f"<!DOCTYPE html><html><body><h3>打开文件失败：{html_lib.escape(str(e))}</h3></body></html>",
            status_code=500,
        )

    with xl:
        for sheet_conf in sheets:
            sheet_name = sheet_conf.get("sheetName", "")
            header_row_conf = sheet_conf.get("headerRow") or {}

            # 支持新格式 {index, data}，也兼容旧格式（列表）
            # data 元素格式：{"index": 1-based列位置, "value": 列名} 或 int（列位置）或 str（列名）
            if isinstance(header_row_conf, dict):
                header_index: Optional[int] = header_row_conf.get("index")
                raw_cols: List = list(header_row_conf.get("data") or [])
                # 将 {"index", "value"} 对象统一展开为 (col_pos_or_name,) 供后续处理
                selected_columns: List = [
                    col["index"] if isinstance(col, dict) else col
                    for col in raw_cols
                ]
            else:
                header_index = None
                selected_columns = list(header_row_conf or [])

            try:
                df = pd.read_excel(xl, sheet_name=sheet_name, header=None)
            except Exception as e:
                logger.error(f"xlsx2html: 读取 sheet={sheet_name} 失败 err={e}")
                parts.append(
                    f"<p>读取 Sheet「{html_lib.escape(sheet_name)}」失败：{html_lib.escape(str(e))}</p>"
                )
                continue

            df = df.where(df.notna(), other=None)
            all_rows = [[_cell_to_str(v) for v in row] for row in df.values.tolist()]

            # 用 index（1-based）直接定位标题行，兼容旧格式按内容匹配
            if header_index is not None:
                header_idx = max(0, header_index - 1)
            else:
                header_idx = 0
                for i, row in enumerate(all_rows):
                    if row == selected_columns:
                        header_idx = i
                        break

            original_headers = all_rows[header_idx] if all_rows else []
            data_rows_raw = all_rows[header_idx + 1:] if header_idx + 1 < len(all_rows) else []

            # 根据 selected_columns 过滤并重排列
            # 元素为 int → 1-based 列位置；为 str → 列名匹配
            if selected_columns:
                col_indices = []
                for col in selected_columns:
                    if isinstance(col, int):
                        idx = col - 1
                        if 0 <= idx < len(original_headers):
                            col_indices.append(idx)
                    else:
                        col_str = str(col)
                        if col_str in original_headers:
                            col_indices.append(original_headers.index(col_str))
                headers = [original_headers[i] for i in col_indices]
                data_rows = [
                    [row[i] if i < len(row) else "" for i in col_indices]
                    for row in data_rows_raw
                ]
            else:
                headers = original_headers
                data_rows = data_rows_raw

            total = len(data_rows)

            caption = f"{display_name} · {sheet_name}"
            style = {
                "header_bg":        "2E75B6",
                "header_color":     "FFFFFF",
                "border_color":     "BFBFBF",
                "font_name":        "仿宋",
                "header_font_size": 10,
                "body_font_size":   9,
                "cell_alignment":   "left",
            }
            full_payload = {
                "caption": caption,
                "headers": headers,
                "rows": data_rows,
                "has_header": True,
                "merge_cells": [],
                "style": style,
            }

            # 用预览数据（前 preview_rows 行）生成 embed 标记 HTML 片段（含 【EMB_xxx】）
            preview_payload = {**full_payload, "rows": data_rows[:preview_rows]}
            spec, snippet = build_embed_marker(
                data=preview_payload,
                embed_type=TYPE_TABLE,
                title=caption,
                record_id=record_id,
                url_builder=lambda eid: f"/doc_editor/embeds/{eid}/go",
            )

            # 入库时替换为完整数据，merge_docx 时写入 DOCX 的是全量行
            spec.payload = full_payload
            try:
                insert_embed_component(spec_to_db_row(spec))
            except Exception as e:
                logger.error(f"xlsx2html: 入库失败 embed_id={spec.embed_id} err={e}")

            parts.append(_render_preview_page(spec, preview_rows))

    escaped_title = html_lib.escape(fileName)
    page = f"""<!DOCTYPE html>
<html lang="zh">
<head>
  <meta charset="UTF-8">
  <title>{escaped_title}</title>
  <style>
    body {{ font-family: "仿宋", serif; padding: 40px; background: #fff; }}
  </style>
</head>
<body>
  {"".join(parts)}
</body>
</html>"""
    return HTMLResponse(content=page)


# ─── 测试接口 ─────────────────────────────────────────────────────────────────

_TEST_PAYLOAD = {
    "caption": "表1 项目人员信息",
    "headers": ["姓名", "部门", "职位", "入职年份"],
    "rows": [
        ["张三", "研发部", "高级工程师", "2019"],
        ["李四", "市场部", "产品经理",  "2021"],
        ["王五", "财务部", "会计",      "2020"],
        ["赵六", "法务部", "法务专员",  "2022"],
    ],
    "has_header": True,
    "col_widths": [3.0, 3.0, 4.0, 3.0],
    "merge_cells": [],
    "style": {
        "header_bg":        "4472C4",
        "header_color":     "FFFFFF",
        "border_color":     "BFBFBF",
        "font_name":        "仿宋",
        "header_font_size": 10,
        "body_font_size":   9,
        "cell_alignment":   "center",
    },
}

# CSV 文件路径（相对于项目根目录）
_CSV_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "static", "3000段表格数据.csv")

# 在 yxdl_embed_components 中标识 CSV 数据的固定标题前缀
_CSV_EMBED_TITLE_PREFIX = "【CSV数据】3000段表格数据"


def _load_csv_payload() -> Dict[str, Any]:
    """从 CSV 文件读取数据，返回 table payload。"""
    headers: List[str] = []
    rows: List[List[str]] = []
    try:
        with open(_CSV_PATH, newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            for i, row in enumerate(reader):
                if i == 0:
                    headers = row
                else:
                    rows.append(row)
    except Exception as e:
        logger.error(f"_load_csv_payload: 读取 CSV 失败 path={_CSV_PATH} err={e}")
    return {
        "caption": "3000段表格数据",
        "headers": headers,
        "rows": rows,
        "has_header": True,
        "merge_cells": [],
        "style": {
            "header_bg":        "2E75B6",
            "header_color":     "FFFFFF",
            "border_color":     "BFBFBF",
            "font_name":        "仿宋",
            "header_font_size": 10,
            "body_font_size":   9,
            "cell_alignment":   "left",
        },
    }


def _get_or_create_csv_embed(record_id: int) -> EmbedSpec:
    """
    查找或创建与 record_id 关联的 CSV 表格 embed。
    若 DB 中已存在则直接返回，否则从 CSV 读取数据入库后返回。
    """
    from psycopg2.extras import RealDictCursor

    # 查找已有的 CSV embed
    sql = """
        SELECT embed_id, embed_type, version, title, display, url, payload,
               record_id, node_id, status
        FROM "yxdl_embed_components"
        WHERE record_id = %s AND embed_type = %s AND title LIKE %s AND status = 1
        ORDER BY create_time ASC LIMIT 1
    """
    with get_db_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(sql, (record_id, TYPE_TABLE, f"{_CSV_EMBED_TITLE_PREFIX}%"))
            row = cur.fetchone()

    if row:
        return spec_from_db_row(dict(row))

    # 未找到，从 CSV 创建
    payload = _load_csv_payload()
    spec, _ = build_embed_marker(
        data=payload,
        embed_type=TYPE_TABLE,
        title=_CSV_EMBED_TITLE_PREFIX,
        record_id=record_id,
        url_builder=lambda eid: f"/doc_editor/embeds/{eid}/go",
    )
    insert_embed_component(spec_to_db_row(spec))
    logger.info(f"_get_or_create_csv_embed: 已入库 embed_id={spec.embed_id} record_id={record_id}")
    return spec


def _render_preview_page(spec: EmbedSpec, preview_rows: int = 10) -> str:
    """渲染包含前 N 行预览表格 + 全部数据链接的 HTML 页面。"""
    from app.utils.embed_table_renderer import render_table_to_html

    payload = spec.payload or {}
    all_rows = payload.get("rows") or []
    total = len(all_rows)

    # 只取前 preview_rows 行用于预览
    preview_payload = {**payload, "rows": all_rows[:preview_rows]}
    preview_spec = EmbedSpec(
        embed_id=spec.embed_id,
        embed_type=spec.embed_type,
        url=spec.url,
        payload=preview_payload,
        version=spec.version,
        title=spec.title,
        display=spec.display,
        record_id=spec.record_id,
        node_id=spec.node_id,
    )
    table_html = render_table_to_html(preview_spec)

    full_url = f"/doc_editor/embeds/test/html/full?embed_id={html_lib.escape(spec.embed_id)}"

    # 【EMB_xxx】锚点段落必须出现在预览内容之前：
    # DOCX 转换服务会将其转为一个独立段落，merge_docx_office_server 靠此文本
    # 定位占位符位置，替换为完整 Word 表格并删除后面的预览表格。
    anchor = (
        f'<a href="{html_lib.escape(spec.url or full_url)}" target="_blank" '
        f'data-embed-anchor="1" data-embed-id="{html_lib.escape(spec.embed_id)}">'
        f'【{html_lib.escape(spec.embed_id)}】</a>'
    )

    page = f"""
  <p style="display:none">{anchor}</p>
  {table_html}
  <p class="tip">
    仅显示前 {preview_rows} 行数据。
    <a href="{full_url}" target="_blank">点击查看全部 {total} 行数据</a>
  </p>
"""
    return page


@router.get(
    "/embeds/test/html/full",
    summary="【测试】全部表格数据页面",
    response_class=HTMLResponse,
)
async def test_html_full(
    embed_id: Optional[str] = Query(None, description="embed_id，与 record_id 二选一"),
    record_id: Optional[int] = Query(None, description="文档 record_id，与 embed_id 二选一"),
):
    """返回完整表格数据的 HTML 页面（通过 embed_id 或 record_id 定位数据）。"""
    from app.utils.embed_table_renderer import render_table_to_html

    spec: Optional[EmbedSpec] = None

    if embed_id:
        row = get_embed_component(embed_id)
        if row:
            spec = spec_from_db_row(row)
    elif record_id:
        try:
            spec = _get_or_create_csv_embed(record_id)
        except Exception as e:
            logger.error(f"test_html_full: 获取 CSV embed 失败 err={e}")

    if spec is None:
        return HTMLResponse(content="<h3>未找到数据，请提供有效的 embed_id 或 record_id</h3>", status_code=404)

    payload = spec.payload or {}
    total = len(payload.get("rows") or [])
    caption = html_lib.escape(payload.get("caption") or "表格数据")
    table_html = render_table_to_html(spec)

    page = f"""<!DOCTYPE html>
<html lang="zh">
<head>
  <meta charset="UTF-8">
  <title>{caption} - 全部数据</title>
  <style>
    body {{ font-family: "仿宋", serif; padding: 40px; background: #fff; color: #222; }}
    .info {{ margin-bottom: 12px; font-size: 13px; color: #555; }}
  </style>
</head>
<body>
  <h3>{caption}（全部 {total} 行）</h3>
  <p class="info" style="display:none">embed_id: {html_lib.escape(spec.embed_id)}</p >
  {table_html}
</body>
</html>"""
    return HTMLResponse(content=page)


@router.get(
    "/embeds/test/html",
    summary="【测试】HTML表格预览渲染（前10行 + 全部数据链接）",
    response_class=HTMLResponse,
)
async def test_html_preview(
    caption: str = Query("表1 项目人员信息", description="表格标题（无 record_id 时生效）"),
    record_id: Optional[int] = Query(None, description="关联文档 record_id；传入后使用 CSV 数据"),
    preview_rows: int = Query(10, ge=1, le=100, description="预览行数，默认10"),
):
    """
    - 传入 record_id：从数据库（或 CSV 首次导入）加载表格数据，展示前 preview_rows 行，
      并在表格下方附「查看全部」链接。
    - 不传 record_id：使用内置示例数据（向后兼容）。
    """
    from app.utils.embed_table_renderer import render_table_to_html

    if record_id is not None:
        try:
            spec = _get_or_create_csv_embed(record_id)
        except Exception as e:
            logger.error(f"test_html_preview: 获取 CSV embed 失败 err={e}")
            return HTMLResponse(content=f"<h3>数据加载失败：{html_lib.escape(str(e))}</h3>", status_code=500)
        return HTMLResponse(content=_render_preview_page(spec, preview_rows))

    # 向后兼容：无 record_id 时使用内置示例数据
    payload = {**_TEST_PAYLOAD, "caption": caption}
    spec, _ = build_embed_marker(data=payload, embed_type=TYPE_TABLE, title=caption)
    table_html = render_table_to_html(spec)
    page = f"""<!DOCTYPE html>
<html lang="zh">
<head>
  <meta charset="UTF-8">
  <title>表格预览</title>
  <style>
    body {{ font-family: "仿宋", serif; padding: 40px; background: #fff; }}
  </style>
</head>
<body>
  {table_html}
</body>
</html>"""
    return HTMLResponse(content=page)


@router.post(
    "/embeds/test/docx",
    summary="【测试】DOCX表格替换渲染，返回可下载的 .docx 文件",
)
async def test_docx_replace(
        payload: Optional[Dict[str, Any]] = Body(
            None,
            description="表格 payload；不传则使用内置示例数据",
        ),
):
    try:
        from docx import Document
    except ImportError:
        return unified_response(500, "python-docx 未安装")

    from app.utils.embed_marker import TYPE_TABLE, VISIBLE_PLACEHOLDER_FMT, build_embed_marker
    from app.utils.embed_table_renderer import render_table_to_docx

    # 👇 关键修复点：如果传进来的是空字典 {}，也会使用默认测试数据
    # 只有当传来的 payload 里面真正包含 "rows" 或 "headers" 时，才使用传入的数据，否则一律使用测试数据
    tbl_payload = payload if (payload and ("rows" in payload or "headers" in payload)) else _TEST_PAYLOAD

    spec, _ = build_embed_marker(
        data=tbl_payload,
        embed_type=TYPE_TABLE,
        title=tbl_payload.get("caption", "测试表格"),
    )

    doc = Document()
    placeholder_text = VISIBLE_PLACEHOLDER_FMT.format(embed_id=spec.embed_id)
    placeholder_para = doc.add_paragraph(placeholder_text)

    try:
        render_table_to_docx(spec, placeholder_para, doc)
    except Exception as e:
        logger.exception("test_docx_replace: 渲染失败")
        return unified_response(500, f"表格渲染失败：{e}")

    import tempfile
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", prefix="embed_test_", delete=False)
    tmp.close()
    doc.save(tmp.name)

    from fastapi.responses import FileResponse
    return FileResponse(
        path=tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"embed_table_test_{spec.embed_id}.docx",
    )